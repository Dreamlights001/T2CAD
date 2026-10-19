"""Export the measured Markdown report to an editable, image-rich Word document."""
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/course-report-measured.md"
OUTPUT = ROOT / "docs/course-report.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text.strip())
    run.bold = bold
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def clean_inline(text):
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    return text.replace("`", "")


def add_paragraph(document, text, style=None):
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(clean_inline(text))
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(10.5)
    paragraph.paragraph_format.line_spacing = 1.35
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def add_table(document, rows):
    values = [[clean_inline(x) for x in row] for row in rows if row]
    if len(values) < 2:
        return
    values = [row for row in values if not all(set(x) <= {"-", ":", " "} for x in row)]
    table = document.add_table(rows=len(values), cols=max(len(row) for row in values))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r, row in enumerate(values):
        for c in range(len(table.columns)):
            set_cell_text(table.cell(r, c), row[c] if c < len(row) else "", bold=r == 0)
            if r == 0:
                set_cell_shading(table.cell(r, c), "DCEDEE")
    document.add_paragraph()


def export():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    normal = document.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)
    for name, size, color in [("Title", 22, "1F5962"), ("Heading 1", 17, "286D76"), ("Heading 2", 13, "397B83"), ("Heading 3", 11, "4C6870")]:
        style = document.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("T2CAD CAD 技术课程实验报告  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = min(len(heading.group(1)), 3)
            paragraph = document.add_paragraph(style="Title" if level == 1 else f"Heading {level}")
            run = paragraph.add_run(clean_inline(heading.group(2)))
            run.font.name = "Noto Sans CJK SC"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
            index += 1
            continue
        image = re.match(r"^!\[([^]]*)\]\(([^)]+)\)$", line)
        if image:
            path = ROOT / "docs" / image.group(2).replace("screenshots/", "screenshots/")
            if path.exists():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(path), width=Inches(6.0))
                caption = document.add_paragraph(image.group(1))
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.runs[0].italic = True
                caption.runs[0].font.size = Pt(9)
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue
        if line.startswith("```"):
            index += 1
            code = []
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            run = paragraph.add_run("\n".join(code))
            run.font.name = "DejaVu Sans Mono"
            run.font.size = Pt(8)
            continue
        if re.match(r"^\d+\.\s+", line):
            add_paragraph(document, re.sub(r"^\d+\.\s+", "", line), "List Number")
        elif line.startswith("- "):
            add_paragraph(document, line[2:], "List Bullet")
        else:
            add_paragraph(document, line)
        index += 1

    document.core_properties.title = "T2CAD 本地大模型驱动的参数化零件生成 CAD 原型系统"
    document.core_properties.subject = "CAD 技术课程实验报告"
    document.core_properties.author = ""
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    export()
